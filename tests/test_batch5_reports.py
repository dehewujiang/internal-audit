#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
[INPUT]: 依赖 program_generator.py、ledger/export.py、
          report-generator/script/progress_report.py 与 report_to_word.py
[OUTPUT]: 断言式测试输出（✅/❌）+ 汇总；exit 0=全过 / 1=有失败
[POS]:    tests/ 下的专项测试，锁死第五拨（汇报包）四件事：
          检查单现有 S/作废可见 / 大桌总览作废一览 / 进度表Excel /
          报告Word版。每条都能被"撤销改动"验红。
[PROTOCOL]: 变更时更新此头部, 然后检查同级 CLAUDE.md
"""

import json
import subprocess
import sys
import tempfile
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

REPO_ROOT = Path(__file__).resolve().parent.parent
PROG_GEN = REPO_ROOT / "internal-audit-program-generator" / "script" / "program_generator.py"
EXPORT = REPO_ROOT / "ledger" / "export.py"
REP_SCRIPT = REPO_ROOT / "internal-audit-report-generator" / "script"
SANDBOX = Path(tempfile.mkdtemp(prefix="batch5_"))

failures = []


def check(label, ok, detail=""):
    print(f"  {'✅' if ok else '❌'} {label}" + (f" —— {detail}" if detail else ""))
    if not ok:
        failures.append(label)


MD = """# 废料管理审计程序

### 2.1 风险识别清单

| 风险编号 | 风险名称 | 风险描述 | 来源标注 |
|----------|----------|----------|----------|
| R01 | 夜班单人称重 | 夜班称重单人操作无复核 | 【推演】 |
| R02 | 地磅年久失修 | 【已删除-原因：地磅已换新】 | 【推演】 |

<!-- track A -->
| 程序编号 | 风险编号 | 测试程序 | 取证方式 | 判定标准 |
|----------|----------|----------|----------|----------|
| A1.1 | R01 | 夜班称重跟访 | 称重单 | 无签字笔数=0 |
| A1.2 | R02 | 地磅精度测试【已删除-原因：地磅已换新】 | 地磅检定证书 | 精度达标 |
<!-- end track A -->

## 十、访谈补充测试程序

| 补充编号 | 风险编号 | 测试程序 | 取证方式 | 判定标准 |
|----------|----------|----------|----------|----------|
| S01 | R01 | 夜班突击复核 | 称重单 | 有双人签字 |
"""


def _ws():
    ws = Path(tempfile.mkdtemp(prefix="b5ws_", dir=str(SANDBOX)))
    iw = ws / "internal-audit-workspace"
    (iw / "audit-programs").mkdir(parents=True)
    (iw / "audit-programs" / "废料管理审计程序_v1.0.md").write_text(MD, encoding="utf-8")
    (iw / "findings").mkdir(exist_ok=True)
    (iw / "findings" / "F-2026-005.json").write_text(json.dumps({
        "finding_id": "F-2026-005", "title": "夜班称重单人无复核",
        "risk_level": "高", "status": "待整改", "responsible": "制造部长",
        "related_procedures": ["A1.1", "S01"],
        "evidence": [{"name": "称重单7张"}]}, ensure_ascii=False), encoding="utf-8")
    (iw / "audit-table").mkdir(exist_ok=True)
    (iw / "audit-table" / "废料.json").write_text(json.dumps({
        "schema_version": "1.2", "table": "废料",
        "left": [{"slot": "确定的毛病", "red": False, "text": "F-2026-005 夜班称重",
                  "ref_finding_ids": ["F-2026-005"]},
                 {"slot": "怀疑偷骗", "red": True, "text": "", "ref_finding_ids": []},
                 {"slot": "说不清的信号", "red": False, "text": "R01 推演风险",
                  "ref_finding_ids": []}],
        "right": [], "drawers": [], "checklist": [], "ingested": {}},
        ensure_ascii=False), encoding="utf-8")
    (iw / "evidence").mkdir(exist_ok=True)
    (iw / "evidence" / "_evidence_catalog.json").write_text(json.dumps({
        "project": "废料", "total_slots": 1, "filled_slots": 1, "items": [
            {"id": "EVD-001", "name": "称重单", "source_track": "A,S",
             "source_programs": ["A1.1", "S01"], "file": "_files/称重单7张.pdf",
             "collected_at": "2026-09-15"}]}, ensure_ascii=False), encoding="utf-8")
    (iw / "reports").mkdir(exist_ok=True)
    return ws


def _sheet_texts(xlsx, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[sheet]
    return ["|".join("" if c.value is None else str(c.value) for c in row)
            for row in ws.iter_rows()]


# ══════════════════════════════════════════════════════════════
def test_program_export_has_s_and_deleted():
    print("\n[1] 检查单Excel：S 加菜与作废章看得见")
    d = Path(tempfile.mkdtemp(prefix="b5p_", dir=str(SANDBOX)))
    md = d / "prog.md"
    md.write_text(MD, encoding="utf-8")
    out = d / "prog.xlsx"
    r = subprocess.run([sys.executable, str(PROG_GEN), str(md), str(out)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    check("exit 0", r.returncode == 0, (r.stdout + r.stderr)[:200])
    from openpyxl import load_workbook
    names = load_workbook(out, read_only=True, data_only=True).sheetnames
    blob = ""
    for n in names:
        blob += "\n".join(_sheet_texts(out, n))
    check("S01 在表里", "S01" in blob, f"sheet={names}")
    check("作废章在表里", "已删除" in blob)


def test_table_export_deleted_sheet():
    print("\n[2] 大桌总览：作废一览")
    ws = _ws()
    t = ws / "internal-audit-workspace" / "audit-table" / "废料.json"
    out1 = ws / "总览.xlsx"
    r = subprocess.run([sys.executable, str(EXPORT), str(t), str(out1)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    check("不带项目也照出三页", r.returncode == 0, (r.stdout + r.stderr)[:150])
    out2 = ws / "总览2.xlsx"
    r = subprocess.run([sys.executable, str(EXPORT), str(t), str(out2),
                        "--workspace", str(ws)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    check("带项目 exit 0", r.returncode == 0, (r.stdout + r.stderr)[:200])
    from openpyxl import load_workbook
    names = load_workbook(out2, read_only=True, data_only=True).sheetnames
    check("多一页作废一览", "作废一览" in names, f"sheet={names}")
    blob = "\n".join(_sheet_texts(out2, "作废一览"))
    check("A1.2 在作废页", "A1.2" in blob, blob[:150])


def test_progress_xlsx():
    print("\n[3] 进度表Excel")
    ws = _ws()
    out = ws / "进度表.xlsx"
    r = subprocess.run(
        [sys.executable, str(REP_SCRIPT / "progress_report.py"),
         "--workspace", str(ws), "--out", str(out)],
        capture_output=True, text=True, encoding="utf-8", errors="replace")
    check("exit 0", r.returncode == 0, (r.stdout + r.stderr)[:300])
    rows = _sheet_texts(out, "风险进度")
    head = rows[0] if rows else ""
    for col in ("风险编号", "哪来的", "查到哪步", "谁手里", "啥状态", "证据", "桌位"):
        check(f"有[{col}]列", col in head, head[:120])
    blob = "\n".join(rows)
    check("R01 行带问题单与桌位", "F-2026-005" in blob and "确定的毛病" in blob)
    check("R02 行标作废", "R02" in blob and "作废" in blob)


def test_word_report():
    print("\n[4] 报告Word版")
    d = Path(tempfile.mkdtemp(prefix="b5w_", dir=str(SANDBOX)))
    md = d / "报告.md"
    md.write_text("# 废料审计报告\n\n## 一、发现\n\n- 夜班单人无复核\n\n| 单号 | 风险 |\n|------|------|\n| F-2026-005 | 高 |\n",
                  encoding="utf-8")
    out = d / "报告.docx"
    r = subprocess.run(
        [sys.executable, str(REP_SCRIPT / "report_to_word.py"),
         str(md), str(out)],
        capture_output=True, text=True, encoding="utf-8", errors="replace")
    check("exit 0", r.returncode == 0, (r.stdout + r.stderr)[:200])
    from docx import Document
    doc = Document(str(out))
    paras = "\n".join(p.text for p in doc.paragraphs)
    check("标题在", "废料审计报告" in paras)
    check("表格在", len(doc.tables) == 1 and "F-2026-005" in doc.tables[0].rows[1].cells[0].text)
    check("条目在", "夜班单人无复核" in paras)


if __name__ == "__main__":
    test_program_export_has_s_and_deleted()
    test_table_export_deleted_sheet()
    test_progress_xlsx()
    test_word_report()
    print(f"\n{'全过' if not failures else f'失败 {len(failures)} 项: {failures}'}")
    sys.exit(1 if failures else 0)
