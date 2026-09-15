#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
report_to_word.py — 审计报告 Markdown 转 Word（给没系统的人看）

[INPUT]:  Markdown 报告文件（标题/表格/条目/引用，原样搬运）
[OUTPUT]: .docx 文件；退出码 0=成功, 2=失败
[POS]:    report-generator/script 的对外汇报零件，只管排版，不管结论对错。
          内容以 MD 为准，不另起新数。
[PROTOCOL]: 变更时更新此头部, 然后检查同级 CLAUDE.md

用法:
    python report_to_word.py 报告.md 报告.docx
"""

import re
import sys
from pathlib import Path


def _is_table_sep(cells) -> bool:
    return bool(cells) and all(re.match(r'^[-:]+$', c) or c == '' for c in cells)


def _split_row(line: str):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def md_to_docx(md_path: str, out_path: str) -> str:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles['Normal']
    style.font.size = Pt(10.5)
    style.font.name = '宋体'

    lines = Path(md_path).read_text(encoding="utf-8-sig").splitlines()
    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.strip().startswith('```'):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            doc.add_paragraph(line)
            i += 1
            continue
        m = re.match(r'^(#{1,3})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue
        if line.strip().startswith('|') and line.strip().endswith('|'):
            cells = _split_row(line)
            # 看下一行是不是分隔行，是才算表
            nxt = _split_row(lines[i + 1]) if i + 1 < len(lines) else []
            if _is_table_sep(nxt):
                body = []
                j = i + 2
                while j < len(lines):
                    lj = lines[j].strip()
                    if not (lj.startswith('|') and lj.endswith('|')):
                        break
                    body.append(_split_row(lj))
                    j += 1
                table = doc.add_table(rows=1 + len(body), cols=len(cells))
                table.style = 'Table Grid'
                for k, h in enumerate(cells):
                    cell = table.rows[0].cells[k]
                    cell.text = h
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.bold = True
                for ri, row in enumerate(body, 1):
                    for k in range(len(cells)):
                        table.rows[ri].cells[k].text = row[k] if k < len(row) else ''
                i = j
                continue
        m = re.match(r'^\s*[-*]\s+(.*)', line)
        if m:
            doc.add_paragraph(m.group(1).strip(), style='List Bullet')
            i += 1
            continue
        m = re.match(r'^\s*>\s?(.*)', line)
        if m:
            doc.add_paragraph(m.group(1).strip())
            i += 1
            continue
        if line.strip() in ('---', '***', '___'):
            i += 1
            continue
        if line.strip():
            doc.add_paragraph(line.strip())
        i += 1

    doc.save(out_path)
    return out_path


def main() -> int:
    if len(sys.argv) != 3:
        print("用法: python report_to_word.py 报告.md 报告.docx")
        return 2
    out = md_to_docx(sys.argv[1], sys.argv[2])
    print(f"Word 报告：{out}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        import traceback
        traceback.print_exc()
        sys.exit(2)
